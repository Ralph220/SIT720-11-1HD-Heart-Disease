"""Build the editable report from verified result files and the user's style template."""
from pathlib import Path
from copy import deepcopy
import json, pandas as pd
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parent;R=ROOT/'results';F=ROOT/'figures'
facts=json.loads((R/'report_facts.json').read_text());stats={r['model']:r for r in facts['unique_summary']};rep=pd.read_csv(R/'reproduction_metrics.csv');pub=pd.read_csv(R/'published_metrics.csv');fm=pd.read_csv(R/'fold_metrics.csv');unc=facts['uncertainty'];base=stats['Baseline stacking'];pro=stats['Proposed pool'];audit=json.loads((R/'audit.json').read_text());split=json.loads((R/'reproduction_split.json').read_text())
d=Document(ROOT/'reports/style_template.docx');protos=[deepcopy(p._p) for p in d.paragraphs]
for el in list(d._element.body):
 if el.tag!=qn('w:sectPr'):d._element.body.remove(el)
# Template fonts, A4 page size, margins and blue heading hierarchy retained.
def p(text,kind=6):
 el=deepcopy(protos[kind]);d._element.body.insert(len(d._element.body)-1,el);pp=Paragraph(el,d._body)
 rp=deepcopy(pp.runs[0]._r.rPr) if pp.runs and pp.runs[0]._r.rPr is not None else None
 for e in list(el):
  if e.tag!=qn('w:pPr'):el.remove(e)
 run=pp.add_run(text)
 if rp is not None:run._r.insert(0,rp)
 pp.paragraph_format.page_break_before=False
 if kind==6:
  pp.paragraph_format.space_after=Pt(7);pp.paragraph_format.line_spacing=1.12;pp.paragraph_format.keep_with_next=False
 return pp
def h(t):return p(t,4)
def sub(t):return p(t,5)
def page(t):
 pp=h(t);pp.paragraph_format.page_break_before=True
def table(rows,widths):
 t=d.add_table(rows=0,cols=len(widths));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for i,row in enumerate(rows):
  cells=t.add_row().cells
  for j,(cell,text,w) in enumerate(zip(cells,row,widths)):
   cell.width=Inches(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;cell.text=str(text)
   pr=cell._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
   pr.append(b)
   if i==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'D9E2F3');pr.append(sh)
   for pp in cell.paragraphs:
    pp.paragraph_format.space_before=Pt(4);pp.paragraph_format.space_after=Pt(4);pp.paragraph_format.line_spacing=1.02
    pp.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or len(str(text))>24 else WD_ALIGN_PARAGRAPH.CENTER
    for rr in pp.runs:rr.font.name='Calibri';rr.font.size=Pt(9.5);rr.bold=i==0
  t.rows[-1]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
 t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 d.add_paragraph().paragraph_format.space_after=Pt(0)
 return t
def fig(name,caption,width=6.5):
 pp=d.add_paragraph();pp.alignment=WD_ALIGN_PARAGRAPH.CENTER;pp.paragraph_format.space_after=Pt(2);pp.add_run().add_picture(str(F/name),width=Inches(width));pp.paragraph_format.keep_with_next=True
 cap=p(caption);cap.paragraph_format.space_after=Pt(8)
 for rr in cap.runs:rr.font.size=Pt(9);rr.italic=True

def n(x):return f'{x:.4f}'
def pct(x):return f'{100*x:.2f}%'
# PAGE 1
p('11.1 High Distinction Task',0);p('Machine Learning Mini Research',1);p('Stacking and Duplicate Aware Heart Disease Classification',2);p('Ralph Henroy Gonsalves  ·  Student ID 226515084',3)
sub('Abstract')
p(f'This study reproduces six classifiers and a stacking ensemble from Bhagat et al. [1] using the cited heart-disease dataset, then evaluates a redesigned probability ensemble. An audit found 723 duplicate rows among 1,025 observations, leaving 302 distinct records. An assumption-based random-row reproduction achieved 100% stacking accuracy, but 202 of 205 test rows had a matching profile in training. Under three repetitions of five-fold evaluation with distinct profiles separated, baseline stacking achieved {pct(base["accuracy_mean"])} mean accuracy and {n(base["brier_mean"])} Brier score. Unique-record training, categorical encoding and regularised probability pooling achieved {pct(pro["accuracy_mean"])} accuracy and {n(pro["brier_mean"])} Brier score. These exploratory findings support a more careful evaluation protocol and a modest improvement within it; they do not establish clinical readiness or reproduce the authors’ exact implementation.')
h('Part 1: Research Problem and Published Methods')
sub('1.1 Problem Definition and Scope')
p('The selected paper investigates whether combining logistic regression, decision tree, random forest, XGBoost, naive Bayes and K-nearest neighbours improves heart-disease classification. The dataset target records the presence or absence of heart disease [1], [2]. It does not specify a future event or prediction horizon, so the paper’s title referring to early heart attack prediction is broader than the outcome verified in this study.')
sub('1.2 Related Work and Research Gap')
p('Ting and Witten [3] support using class probabilities rather than only hard predictions as inputs to a higher-level classifier. XGBoost [4] provides an established tree-boosting system and is one of the paper’s six components. Bhagat et al. [1] report 98.53% stacking accuracy, exceeding their best individual classifiers at 92.68%. That is a result within their study, not proof of current state-of-the-art performance across clinical populations.')
p('Kapoor and Narayanan [5] identify leakage as a source of overoptimistic ML-based scientific findings. Here, the concrete research gap is whether a high reported score is robust to repeated-record handling, explicit feature representation and probability-focused ensemble training. The proposed contribution is an empirical redesign of this pipeline. It uses established methods and makes no claim to invent a universally novel learning algorithm.')
# PAGE 2
page('1.3 Dataset Audit and Reproduction Decisions')
p('The unmodified heart.csv was downloaded from the Kaggle dataset cited in Section 4 of [1]. All 13 predictors were retained. Distinct records were determined using the complete predictor profile; no conflicting labels were found for identical profiles. A profile identifier is not a verified patient identifier.')
table([['Audit item','Observed value'],['Rows and columns','1,025 rows; 13 predictors plus target'],['Classes in original rows','0: 499; 1: 526'],['Distinct records','302; class 0: 138; class 1: 164'],['Extra duplicate rows','723 (70.54% of rows)'],['Missing cells / conflicting profiles','0 / 0'],['Unexpected category ranges','ca includes 4; thal includes 0']],[2.5,4.15])
p('The unexpected category values fall outside the ranges in the paper’s Table 9. Their meaning was not independently established, so they were retained as observed categories rather than silently removed or recoded. No oversampling was used. Section 3.1 describes regression-based imputation, but no imputation was necessary because the downloaded file has no missing cells.')
sub('1.4 Confirmed Details and Assumptions')
p('The same dataset source, predictor set, six model families, five-fold stacking concept and five required evaluation metrics were used. The paper does not give a complete executable parameter specification. The following choices are assumptions, not settings verified from the authors: stratified 80/20 row holdout, seed 42, standard scaling, retaining numeric category codes for the reproduction, probability inputs and a logistic-regression meta-classifier. No seed was selected to match the published score.')
table([['Component','Implementation used'],['Logistic regression','C = 1; max_iter = 2,000; default L2 regularisation'],['Decision tree','Gini criterion; unlimited depth; otherwise library defaults'],['Random forest','200 trees; otherwise library defaults; one worker'],['XGBoost','150 trees; depth 3; learning rate 0.05; subsample and column fraction 1; L2 = 1; histogram tree method'],['Naive Bayes and KNN','GaussianNB defaults; KNN with k = 5 and uniform weights'],['Stacking','Five-fold out-of-fold probability inputs; logistic meta-learner C = 1; no original-feature passthrough']],[1.6,5.05])
p('Scaling was fitted within every base-model training fold; test data were only transformed. This follows the separation principle described in [6]. The held-out reproduction used 820 training and 205 test rows. The five-fold meta-training procedure used only those 820 training rows, but did not remove duplicate profiles in this diagnostic reproduction. Defaults and versions are recorded in the code package.')
# PAGE 3
page('1.5 Published and Reproduced Results')
p('Table 1. Published values transcribed from Table 11 of [1]. Metrics are proportions; precision, recall and F1 below retain the source’s labels.')
heads=['Model','Accuracy','Precision','Recall','F1','AUC'];cols=[1.45,1.04,1.04,1.04,1.04,1.04]
table([heads]+[[r.model]+[n(getattr(r,m)) for m in ['accuracy','precision','recall','f1','auc']] for r in pub.itertuples()],cols)
p('Table 2. Measured results for the assumption-based 80/20 random-row reproduction. The target-positive class is 1; classification threshold is 0.5. AUC uses probabilities.')
table([heads]+[[r.model]+[n(getattr(r,m)) for m in ['accuracy','precision','recall','f1','auc']] for r in rep.itertuples()],cols)
sub('1.6 Interpretation of Discrepancies')
p('Stacking and random forest classified all 205 held-out rows correctly, but 202 test rows had an identical predictor profile in training. The apparent 100% accuracy therefore measures performance largely on repeated profiles. It is not evidence of generalisation to new patients, nor proof that the original authors used the same split. The 1.47 percentage-point difference from the paper’s stacking score cannot be interpreted as a superior method.')
p('The other accuracy differences, reproduced minus published, were LR −3.41, DT +5.86, RF +7.32, XGBoost +4.88, NB −1.46 and KNN +0.49 percentage points. Split membership, assumed hyperparameters, preprocessing and software versions can all contribute. The experiment quantifies these differences without claiming to isolate every cause.')
p('The paper also reports stacking recall as 0.9727 but sensitivity as 0.9853. For the same binary positive class those measures should agree. Its MCC value of 1.0 is also inconsistent with imperfect accuracy when both classes are present. The report therefore uses independently computed metrics and treats the source values as reporting inconsistencies requiring clarification, rather than silently accepting them.')
# PAGE 4
page('Part 2: Proposed Method and Experimental Design')
sub('2.1 Motivation and Method')
p('The audit motivates an ensemble trained on distinct profiles, with categorical features represented explicitly and the combination rule fitted for probability quality. The method retains all six original classifier families. It changes training-record weighting, feature representation and the ensemble objective, rather than merely substituting a classifier or adjusting a single parameter.')
p('Five variables are treated as continuous: age, trestbps, chol, thalach and oldpeak. The other eight predictors are one-hot encoded, with unseen levels ignored at transformation time. Continuous scaling and category discovery are fitted inside each training fold. Original feature values are preserved; no target-informed feature selection is performed.')
p('For each outer training set, duplicate profiles are reduced to one record. Five-fold out-of-fold predictions form a matrix with one probability column for each base learner. Six non-negative weights, constrained to sum to one, minimise the mean squared difference between the weighted probability and the target, plus 0.01 times the sum of squared deviations of the weights from 1/6. The penalty shrinks weights towards equal pooling. It was fixed before the full experiments, not selected using outer test scores.')
p('The optimiser is SLSQP with tolerance 10⁻¹² and a maximum of 1,000 iterations. All fits converged. Base learners are refitted on the complete outer training set, then their test probabilities are combined with the learned weights. This is a regularised convex probability pool; no separate post-hoc calibrator was fitted. Brier score measures overall probability quality, not calibration alone [7].')
sub('2.2 Controlled Ablations')
table([['Variant','Training and ensemble changes'],['Baseline stacking','Keep repeated training rows; numeric category codes; group-separated inner folds; logistic meta-learner'],['Unique stacking','One training row per profile; otherwise the same numeric-code pipeline and logistic meta-learner'],['Typed stacking','Unique training plus scaled continuous and one-hot categorical features; logistic meta-learner'],['Proposed pool','Same unique typed base predictions; replace logistic meta-learner with regularised probability pooling']],[1.6,5.05])
sub('2.3 Validation and Comparison Rules')
p('The 302 distinct profiles are split using stratified five-fold validation repeated with seeds 42, 1042 and 2042. Original duplicate rows are mapped to their profile’s fold. Every model is evaluated on exactly the same test profiles, and all baseline inner splits keep duplicate groups together. Primary scoring gives one vote per test profile; a secondary analysis restores each profile’s original row count as its weight. The six standalone baseline models are also evaluated on these shared outer folds.')
p('Mean Brier score is the primary comparison, with accuracy, precision, recall, F1, AUC and log loss retained as secondary measures. Each model has 15 outer-fold scores. These folds are dependent because training sets overlap. The design was fixed after the earlier LR/RF pilot, so the study remains exploratory and has no untouched external test set.')
# PAGE 5
page('2.4 Performance on Unseen Profiles')
p('Table 3. Mean metrics across 15 shared outer folds, giving each distinct test profile equal weight. Higher is better for the five classification metrics; lower Brier is better. All entries are measured results.')
heads=['Model','Acc.','Prec.','Recall','F1','AUC','Brier'];widths=[1.5,.86,.86,.86,.86,.86,.86]
table([heads]+[[name]+[n(stats[name][m+'_mean']) for m in ['accuracy','precision','recall','f1','auc','brier']] for name in ['LR','DT','RF','XGB','NB','KNN','Baseline stacking','Unique stacking','Typed stacking','Proposed pool']],widths)
fig('accuracy_comparison.png','Figure 1. Mean accuracy with fold standard deviation. Error bars describe variation across folds, not independent-sample confidence intervals.',6.0)
p(f'Baseline stacking averaged {pct(base["accuracy_mean"])} accuracy (SD {100*base["accuracy_std"]:.2f} percentage points). The proposed pool averaged {pct(pro["accuracy_mean"])} (SD {100*pro["accuracy_std"]:.2f}), an increase of {100*(pro["accuracy_mean"]-base["accuracy_mean"]):.2f} percentage points. Mean AUC increased from {n(base["auc_mean"])} to {n(pro["auc_mean"])}. The baseline no longer outperformed every individual classifier once evaluation excluded repeated profiles across folds.')
p('The improvement is relative to the baseline on the same rigorous protocol. The proposed result is lower than the paper’s 98.53%, but those numbers use different test designs and cannot be ranked as a fair head-to-head comparison. The change from the diagnostic reproduction to grouped evaluation also changes the scoring unit and split membership, so it does not isolate a causal leakage effect by itself.')
# PAGE 6
page('2.5 Probability Quality and Ablation Findings')
table([['Variant','Accuracy','Brier mean','Brier SD'],*[ [name,pct(stats[name]['accuracy_mean']),n(stats[name]['brier_mean']),n(stats[name]['brier_std'])] for name in ['Baseline stacking','Unique stacking','Typed stacking','Proposed pool']]],[2.55,1.25,1.45,1.4])
p(f'Removing duplicate training weight reduced mean Brier score from {n(base["brier_mean"])} to {n(stats["Unique stacking"]["brier_mean"])}. Explicit categorical representation reduced it to {n(stats["Typed stacking"]["brier_mean"])}. Replacing the logistic meta-learner with the probability pool gave a further, smaller reduction to {n(pro["brier_mean"])}. Most of the observed change came before the final pooling step; the results do not support attributing the entire gain to the ensemble objective.')
p('The duplicate-removal ablation also changes the inner fold construction from grouped repeated rows to stratified unique rows. That comparison therefore bundles weighting and inner-partition effects. Typed stacking and the proposed pool share exactly the same base probabilities, making their final combination-rule comparison more directly controlled.')
fig('roc_reliability.png','Figure 2. Cross-fitted predictions for repeat 1 only (302 distinct profiles). ROC AUC here is pooled within that repeat and differs from the mean fold AUC in Table 3. Reliability curves use six quantile bins.',6.5)
p('Mean pool weights were LR 0.473, DT 0.023, RF 0.173, XGBoost 0.051, NB 0.067 and KNN 0.214. These are ensemble weights, not feature importance or causal effects. The high weight on logistic regression and low weight on a single tree are consistent with favouring stable probability estimates in a small sample, but the weights alone cannot prove the mechanism.')
p(f'A profile-level bootstrap of paired squared-error differences, retaining all three repeats together, gave a mean proposed-minus-baseline difference of {unc["mean_profile_loss_difference"]:.4f}, with a conditional 95% interval [{unc["conditional_bootstrap_95_interval"][0]:.4f}, {unc["conditional_bootstrap_95_interval"][1]:.4f}] from 5,000 resamples. This describes uncertainty conditional on fitted cross-validated predictions. It omits model-refitting uncertainty and does not make overlapping folds independent; it is exploratory evidence rather than a definitive significance test.')
p('With original row frequencies restored as evaluation weights, baseline accuracy/Brier were 81.60%/0.1329 and proposed accuracy/Brier were 84.53%/0.1152. The direction of the result is therefore robust to this scoring choice. A lower Brier score does not establish better calibration alone; the reliability diagram is descriptive and has no binwise uncertainty bands [7].')
# PAGE 7
page('2.6 Error Analysis and Limitations')
fig('confusion_matrices.png','Figure 3. Repeat 1 confusion matrices at threshold 0.5. Each matrix represents the same 302 distinct profiles, each tested once.',6.0)
p('In repeat 1, baseline stacking made 28 false positives and 25 false negatives. The proposed pool made 26 false positives and 19 false negatives. It corrected 13 baseline mistakes while introducing five new mistakes, giving eight fewer errors overall. This illustrates a trade-off rather than uniformly better predictions for every record. No clinical cost function was supplied, so the threshold stayed at 0.5 and no claim is made that it is clinically optimal.')
sub('2.7 Strengths and Practical Interpretation')
p('The study implements every classifier family in the selected paper, preserves its feature set for the reproduction, saves every test prediction and compares the extension against a common baseline under shared splits. It reports both successes and remaining errors. The pipeline design addresses a verified property of the downloaded data, while keeping published benchmarks separate from new measurements.')
p('The comparison also separates classification accuracy from probability quality. Guo et al. [8] demonstrate this distinction in neural-network settings; their specific temperature-scaling result is not tested here. This study instead evaluates a probability-pooling objective on tabular data. It does not claim that the pool is a clinically calibrated risk score.')
sub('2.8 Limitations and Future Work')
p('Only 302 distinct predictor profiles are available, and patient identifiers are absent. Identical profiles may occasionally represent distinct patients, while non-identical profiles may still belong to the same patient. Grouping exact profiles prevents measured duplicate overlap but cannot guarantee patient-level independence. Dataset provenance, category semantics and representativeness remain incompletely verified.')
p('The original implementation is underspecified, so this is an assumption-based reproduction rather than exact replication. The fixed hyperparameters were not exhaustively optimised. The proposed method was designed after an exploratory pilot, and repeated cross-validation on the same dataset cannot remove all analyst-selection effects. There is no independent external validation, temporal evaluation, subgroup fairness assessment or clinical utility analysis.')
p('Future work should verify the source records and patient identities, pre-specify the method on a genuinely external dataset, and quantify uncertainty using a design that refits the full pipeline. Further comparisons could distinguish unique-record weighting from inner-fold construction, investigate category coding, and test whether the modest gain from probability pooling persists across samples. Clinical deployment is outside the evidence established here.')
# PAGE 8
page('Part 3: Reproducibility and Presentation')
sub('3.1 Reproduction Instructions')
p('The accompanying project includes the unchanged dataset, selected paper, protocol, Python source, executed notebook, every prediction, metrics, split indices, figures and this report. Python 3.12.14, NumPy 2.3.5, pandas 2.2.3, SciPy 1.17.0, scikit-learn 1.8.0 and xgboost-cpu 3.0.5 were used. Run the following from the extracted project folder in an environment with the pinned requirements:')
for command in ['python -m pip install -r requirements.txt','python run_research.py','python analyse_results.py','python build_report.py']:
 pp=p(command)
 for rr in pp.runs:rr.font.name='Consolas';rr.font.size=Pt(9.5)
p('run_research.py generates the experiment results; analyse_results.py independently checks confusion-count metrics and regenerates figures; build_report.py regenerates the editable document from the saved results. The report PDF was exported from that document. The README gives notebook and PDF export instructions. The data SHA-256 is recorded in protocol.json and checked before fitting.')
p('A second full run in a separate directory with no prior result files reproduced the checked metrics, predictions, split assignments, learned weights and bootstrap output byte for byte. All 300 model/fold/weighting records were independently checked for accuracy, precision, recall, F1 and Brier score. This verifies consistency of this implementation, not correctness of unobserved details in the original authors’ code.')
sub('3.2 Video and Source Links')
p('Video presentation link: TO BE ADDED after recording. The presentation must show the student’s face, screen sharing and narration, last no more than six minutes, and demonstrate selected code execution, comparisons and limitations. A timed outline is included in VIDEO_OUTLINE.md.')
p('Accessible code and data link: TO BE ADDED after uploading the project ZIP to OneDrive/Dropbox or an accessible GitHub repository. A local file link alone is not an independently accessible submission link. Check access permissions from a separate browser session before submission.')
sub('3.3 Conclusion')
p(f'The diagnostic reproduction obtained perfect random-row stacking accuracy, but nearly all test rows shared a profile with training data. On a shared evaluation of unseen profiles, the redesigned pipeline improved mean accuracy by {100*(pro["accuracy_mean"]-base["accuracy_mean"]):.2f} percentage points and reduced Brier score from {n(base["brier_mean"])} to {n(pro["brier_mean"])}. The empirical contribution is a more transparent evaluation and a modest improvement within that evaluation. The evidence does not support claims of near-perfect new-patient prediction, universal superiority or clinical readiness.')
sub('Acknowledgement of Generative AI Use')
p('ChatGPT/Codex generated and executed the Python implementation, assisted with the dataset audit and experimental design, produced the figures and drafted this report using the supplied template. The numerical results came from executed code rather than invented values. Automated checks and a repeat run verified output consistency. Student verification, independent interpretation and the presentation have not been completed by the assistant and are not asserted here.')
p('The task brief permits GenAI for planning, brainstorming and grammar/spelling editing. This package includes more extensive AI assistance. It is a research reference for student review and must not be represented as independently authored work or submitted unchanged without resolving that use with the unit’s requirements. The student remains responsible for an accurate final acknowledgement and demonstrated understanding.')
# PAGE 9
page('References')
refs=[
'[1] M. Bhagat, A. Sharma, and P. Agarwal, “An efficient stacking-based ensemble technique for early heart attack prediction,” Multimedia Tools and Applications, vol. 84, pp. 36351–36375, 2025, doi: 10.1007/s11042-024-19293-7. Published online Apr. 30, 2024.',
'[2] johnsmith88, “Heart Disease Dataset,” Kaggle. Accessed Sep. 21, 2026. [Online]. Available: https://www.kaggle.com/datasets/johnsmith88/heart-disease-dataset',
'[3] K. M. Ting and I. H. Witten, “Issues in stacked generalization,” Journal of Artificial Intelligence Research, vol. 10, pp. 271–289, 1999, doi: 10.1613/jair.594. https://jair.org/index.php/jair/article/view/10228',
'[4] T. Chen and C. Guestrin, “XGBoost: A scalable tree boosting system,” 2016, arXiv:1603.02754. [Online]. Available: https://arxiv.org/abs/1603.02754',
'[5] S. Kapoor and A. Narayanan, “Leakage and the reproducibility crisis in ML-based science,” 2022, arXiv:2207.07048. [Online]. Available: https://arxiv.org/abs/2207.07048',
'[6] scikit-learn developers, “Common pitfalls and recommended practices,” scikit-learn documentation. Accessed Sep. 21, 2026. [Online]. Available: https://scikit-learn.org/stable/common_pitfalls.html',
'[7] scikit-learn developers, “Probability calibration,” scikit-learn documentation. Accessed Sep. 21, 2026. [Online]. Available: https://scikit-learn.org/stable/modules/calibration.html',
'[8] C. Guo, G. Pleiss, Y. Sun, and K. Q. Weinberger, “On calibration of modern neural networks,” in Proc. 34th Int. Conf. Machine Learning, PMLR, vol. 70, pp. 1321–1330, 2017. [Online]. Available: https://proceedings.mlr.press/v70/guo17a.html']
for ref in refs:p(ref)
sub('Appendix A: Metric Definitions and Evidence Files')
p('Accuracy is the fraction of correct labels. Precision is TP/(TP + FP), recall is TP/(TP + FN), and F1 is 2TP/(2TP + FP + FN), with zero returned if a denominator is zero. ROC AUC assesses probability ranking. Binary Brier score is the average squared difference between predicted probability and the 0/1 target. Log loss uses clipped probabilities through scikit-learn. Primary tables average fold metrics rather than pooling all repeat predictions.')
table([['Evidence','File in results folder'],['Published and reproduced metrics','published_metrics.csv; reproduction_metrics.csv'],['All outer-fold scores and predictions','fold_metrics.csv; outer_predictions.csv'],['Split assignments and learned weights','outer_splits.json; pool_weights.csv'],['Audit and source checks','audit.json; environment.json; run_info.json'],['Error analysis and uncertainty','error_analysis_repeat1.csv; uncertainty.json'],['Validation of reported outputs','verification.txt; clean_rerun_verification.json']],[2.6,4.05])
p('The earlier two-model pilot used a different outer-split construction and is not mixed into the full-study result tables. All final numerical comparisons above come from the fixed protocol in the accompanying project. Figures 2 and 3 intentionally use only repeat 1; other tables use all 15 folds.')
# Footer only page numbering, no repeated name/date.
for sec in d.sections:
 for pp in sec.footer.paragraphs:pp.clear()
 pp=sec.footer.paragraphs[0];pp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');pp._p.append(field)
d.core_properties.title='11.1 High Distinction Machine Learning Research';d.core_properties.author='Ralph Henroy Gonsalves';d.core_properties.subject='Heart disease classification research'
path=ROOT/'reports/SIT720_11_1HD_Research_Report.docx';d.save(path);print(path)
